# utils/file_parser.py
import os
import PyPDF2
from pptx import Presentation
import pandas as pd
from xmindparser import xmind_to_dict
from utils.logger import logger


class FileParser:

    @classmethod
    def parse(cls, file_path: str, max_chars: int = 500000, original_filename: str = None) -> str:
        """
        统一解析入口
        :param file_path: 本地临时文件路径
        :param max_chars: 字符截断阈值
        :param original_filename: [新增] 原始文件名，用于在 Prompt 中展示给 AI
        """
        if not os.path.exists(file_path):
            return ""

        ext = os.path.splitext(file_path)[1].lower()
        # 优先使用传入的原始文件名，如果没有则回退到本地文件名
        display_name = original_filename if original_filename else os.path.basename(file_path)
        content = ""

        try:
            if ext == '.pdf':
                content = cls._parse_pdf(file_path, max_chars)
            elif ext == '.pptx':
                content = cls._parse_pptx(file_path, max_chars)
            elif ext in ['.xlsx', '.xls']:
                content = cls._parse_excel(file_path)
            elif ext == '.xmind':
                content = cls._parse_xmind(file_path)
            elif ext in ['.txt', '.md', '.py', '.java', '.json', '.csv', '.xml', '.html', '.css', '.js', '.sql', '.sh',
                         '.c', '.cpp']:
                with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                    content = f.read(max_chars)
            else:
                return f"[提示: 不支持解析 {ext} 格式的文件: {display_name}]"

        except Exception as e:
            error_msg = f"文件解析失败 {display_name}: {str(e)}"
            logger.error(error_msg)
            return f"[系统提示: {error_msg}]"

        # 截断检查
        if len(content) > max_chars:
            content = content[:max_chars] + "\n...[单文件内容过长已截断]..."

        if content:
            # 这里在 Prompt 中明确标识文件名，方便 AI 区分
            return f"\n\n=== 开始文件内容: {display_name} ===\n{content}\n=== 结束文件内容 ===\n"
        return ""

    # ... (下方的 _parse_pdf 等静态方法保持不变，无需修改) ...
    # 只需要确保 parse 方法头部更新即可
    @staticmethod
    def _parse_pdf(path, limit):
        # ... (保持原样)
        text = ""
        with open(path, 'rb') as f:
            reader = PyPDF2.PdfReader(f)
            for i, page in enumerate(reader.pages):
                extracted = page.extract_text()
                if extracted:
                    text += f"\n[Page {i + 1}]\n{extracted}"
                if len(text) > limit: break
        return text

    @staticmethod
    def _parse_pptx(path, limit):
        # ... (保持原样)
        text = ""
        prs = Presentation(path)
        for i, slide in enumerate(prs.slides):
            slide_text = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    slide_text.append(shape.text)
            if slide_text:
                text += f"\n[Slide {i + 1}]\n" + "\n".join(slide_text)
            if len(text) > limit: break
        return text

    @staticmethod
    def _parse_excel(path):
        # ... (保持原样)
        text = ""
        dfs = pd.read_excel(path, sheet_name=None)
        for sheet_name, df in dfs.items():
            if len(df) > 1000:
                df = df.head(1000)
                text += f"\n\n### Sheet: {sheet_name} (仅展示前1000行)\n"
            else:
                text += f"\n\n### Sheet: {sheet_name}\n"
            text += df.to_markdown(index=False)
        return text

    @staticmethod
    def _parse_xmind(path):
        # ... (保持原样)
        text = ""
        sheets = xmind_to_dict(path)
        for sheet in sheets:
            root_topic = sheet.get('topic', {})
            text += f"\n\n### 画布: {sheet.get('title', '未命名')}\n"
            text += FileParser._xmind_recursion(root_topic)
        return text

    @staticmethod
    def _xmind_recursion(data, level=0):
        # ... (保持原样)
        indent = "  " * level
        title = data.get('title', '无标题')
        result = f"{indent}- {title}\n"
        for topic in data.get('topics', []):
            result += FileParser._xmind_recursion(topic, level + 1)
        return result